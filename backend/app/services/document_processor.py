"""
LegalLens — Document Processor
Extracts text from PDF and DOCX files with page/section metadata.
"""

import os
import re
from pathlib import Path
from typing import Optional

import pymupdf as fitz  # PyMuPDF
from docx import Document as DocxDocument


from app.utils.helpers import clean_extracted_text


class DocumentPage:
    """Represents a single page of extracted text with metadata."""

    def __init__(self, page_number: int, text: str, section_header: Optional[str] = None):
        self.page_number = page_number
        self.text = text
        self.section_header = section_header

    def to_dict(self) -> dict:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "section_header": self.section_header,
        }


class ExtractedDocument:
    """Container for all extracted document data."""

    def __init__(self, filename: str, file_type: str):
        self.filename = filename
        self.file_type = file_type
        self.pages: list[DocumentPage] = []
        self.full_text: str = ""
        self.page_count: int = 0
        self.is_scanned: bool = False

    @property
    def total_characters(self) -> int:
        return len(self.full_text)

    @property
    def total_words(self) -> int:
        return len(self.full_text.split())


class DocumentProcessor:
    """Handles text extraction from PDF and DOCX files."""

    SUPPORTED_EXTENSIONS = {"pdf", "docx"}

    @staticmethod
    def validate_file(filename: str, file_size: int, max_size_mb: int = 50) -> tuple[bool, str]:
        """Validate file type and size.

        Returns:
            Tuple of (is_valid, error_message).
        """
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if ext not in DocumentProcessor.SUPPORTED_EXTENSIONS:
            return False, f"Unsupported file type: .{ext}. Supported: .pdf, .docx"

        max_bytes = max_size_mb * 1024 * 1024
        if file_size > max_bytes:
            return False, f"File too large. Maximum size: {max_size_mb}MB"

        return True, ""

    @staticmethod
    async def extract_from_pdf(file_path: str) -> ExtractedDocument:
        """Extract text from a PDF file using PyMuPDF with OCR fallback.

        Preserves page boundaries and detects section headers.
        If a page contains scanned images with no text layer, triggers
        Gemini Multimodal OCR transcription.
        """
        from app.services.llm_service import llm_service

        filename = os.path.basename(file_path)
        doc = ExtractedDocument(filename=filename, file_type="pdf")

        pdf = fitz.open(file_path)
        try:
            # Reject encrypted / password-protected PDFs
            if pdf.is_encrypted:
                raise ValueError(
                    "This PDF is password-protected. Please provide an unlocked copy."
                )

            doc.page_count = len(pdf)
            all_text_parts = []

            for page_num in range(len(pdf)):
                page = pdf[page_num]
                text = page.get_text("text")
                text = clean_extracted_text(text)

                # Scanned page detection (< 30 characters on a page)
                if len(text.strip()) < 30:
                    try:
                        # Render high-resolution pixmap of the page
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        ocr_text = await llm_service.transcribe_image(img_bytes, mime_type="image/png")
                        if ocr_text:
                            text = clean_extracted_text(ocr_text)
                            doc.is_scanned = True
                        elif not text.strip():
                            text = (
                                f"[Scanned page {page_num + 1}: Image detected without direct text layer. "
                                "Live Gemini API key required for full Multimodal OCR.]"
                            )
                            doc.is_scanned = True
                    except Exception:
                        pass

                # Try to detect section headers
                section_header = DocumentProcessor._detect_section_header(text)

                doc.pages.append(DocumentPage(
                    page_number=page_num + 1,
                    text=text,
                    section_header=section_header,
                ))
                all_text_parts.append(text)

            doc.full_text = "\n\n".join(all_text_parts)
        finally:
            pdf.close()

        return doc

    @staticmethod
    def extract_from_docx(file_path: str) -> ExtractedDocument:
        """Extract text from a DOCX file using python-docx.

        Preserves paragraph structure, heading hierarchy, and table content.
        """
        filename = os.path.basename(file_path)
        doc = ExtractedDocument(filename=filename, file_type="docx")

        docx = DocxDocument(file_path)

        # Collect all content elements (paragraphs and tables) in document order.
        # python-docx exposes element.body which contains both <w:p> and <w:tbl>
        # in document order, so we iterate the XML children to preserve ordering.
        content_items: list[str] = []
        current_section: Optional[str] = None

        body = docx.element.body
        # Namespace for Word XML
        _wp = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

        for child in body:
            tag = child.tag
            if tag == f'{_wp}p':  # paragraph
                # Find the matching paragraph object
                for para in docx.paragraphs:
                    if para._element is child:
                        text = para.text.strip()
                        if not text:
                            break
                        if para.style and para.style.name and para.style.name.startswith("Heading"):
                            current_section = text
                        content_items.append(text)
                        break
            elif tag == f'{_wp}tbl':  # table
                for table in docx.tables:
                    if table._element is child:
                        table_text = DocumentProcessor._extract_table_text(table)
                        if table_text.strip():
                            content_items.append(table_text)
                        break

        # Group content into logical pages (~3000 chars each for chunking)
        current_page_text: list[str] = []
        current_page_num = 1
        char_count = 0
        page_size = 3000

        for item in content_items:
            current_page_text.append(item)
            char_count += len(item)

            if char_count >= page_size:
                page_text = "\n".join(current_page_text)
                doc.pages.append(DocumentPage(
                    page_number=current_page_num,
                    text=clean_extracted_text(page_text),
                    section_header=current_section,
                ))
                current_page_text = []
                char_count = 0
                current_page_num += 1

        # Add remaining text
        if current_page_text:
            page_text = "\n".join(current_page_text)
            doc.pages.append(DocumentPage(
                page_number=current_page_num,
                text=clean_extracted_text(page_text),
                section_header=current_section,
            ))

        doc.page_count = len(doc.pages) or 1
        doc.full_text = "\n\n".join(p.text for p in doc.pages)
        return doc

    @staticmethod
    def _extract_table_text(table) -> str:
        """Extract text from a DOCX table into a readable string."""
        rows_text = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_text.append(" | ".join(cells))
        return "\n".join(rows_text)

    @staticmethod
    async def extract(file_path: str) -> ExtractedDocument:
        """Extract text from a file based on its extension.

        Args:
            file_path: Path to the document file.

        Returns:
            ExtractedDocument with all extracted text and metadata.

        Raises:
            ValueError: If the file type is not supported.
        """
        import asyncio

        ext = file_path.rsplit(".", 1)[-1].lower()

        if ext == "pdf":
            return await DocumentProcessor.extract_from_pdf(file_path)
        elif ext == "docx":
            # extract_from_docx is synchronous (python-docx / XML I/O); run it in a
            # thread-pool executor so it doesn't block the asyncio event loop.
            loop = asyncio.get_event_loop()
            return await loop.run_in_executor(
                None, DocumentProcessor.extract_from_docx, file_path
            )
        else:
            raise ValueError(f"Unsupported file type: .{ext}")

    @staticmethod
    def _detect_section_header(text: str) -> Optional[str]:
        """Attempt to detect a section header from page text.

        Looks for patterns like:
        - "SECTION 1: Title"
        - "1. Title"
        - "Article I"
        - "CLAUSE 1"
        """
        lines = text.strip().split("\n")[:5]  # Check first 5 lines

        for line in lines:
            line = line.strip()
            if not line:
                continue

            # Match numbered sections/clauses
            patterns = [
                r'^(?:SECTION|Section|CLAUSE|Clause|Article|ARTICLE)\s+\d+',
                r'^\d+\.\s+[A-Z]',
                r'^[IVXLC]+\.\s+',
                r'^(?:SCHEDULE|Schedule|APPENDIX|Appendix|ANNEXURE|Annexure)\s+',
            ]

            for pattern in patterns:
                if re.match(pattern, line):
                    return line[:100]  # Truncate long headers

        return None
